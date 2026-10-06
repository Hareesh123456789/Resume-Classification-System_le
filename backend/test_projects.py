"""Run project checks, then optionally analyze a PDF/DOCX resume."""

import argparse
import json
from pathlib import Path

from extractor.project_extractor import extract_projects


REPORTED_TEXT = """CAREER OBJECTIVE
Seeking an internship opportunity in AI and ML.
SKILLS
Python, HTML, Oracle
EDUCATION
B.Tech. (AI&ML)
2027
PROJECT EXPERIENCE
AI-Based Deepfake Image Detection Using Transfer Learning
1. Deepfake Image Detection.
2.High Detection Accuracy.
CERTIFICATIONS
1) Introduction to Machine learning (12 weeks)
2) Introduction to Artificial intelligence (12 weeks)
Aditya institute of technology and
management, Tekkali
3.Fast and Efficient Prediction.
DECLARATION
The above information is correct.
"""

REPORTED_EXPECTED = [{
    "title": "AI-Based Deepfake Image Detection Using Transfer Learning",
    "technologies": [],
    "description": "Deepfake Image Detection.\nHigh Detection Accuracy.",
}]


def check_examples():
    cases = [
        (
            "reported PDF text and section boundary",
            REPORTED_TEXT,
            REPORTED_EXPECTED,
        ),
        ("multiple labeled projects", """PROJECTS
Project: Resume Analyzer
Technologies: Python, Flask, SQLite
Description: Parsed resumes and ranked applicants.
Project Title: Shopping App
Tech Stack: Flutter, Firebase
- Built login and shopping screens.
EDUCATION
B.Tech 2027
""", [
            {
                "title": "Resume Analyzer",
                "technologies": ["Python", "Flask", "SQLite"],
                "description": "Parsed resumes and ranked applicants.",
            },
            {
                "title": "Shopping App",
                "technologies": ["Flutter", "Firebase"],
                "description": "Built login and shopping screens.",
            },
        ]),
        ("numbered titles and description technologies", """1. Resume Analyzer
- Built an API using Python and Flask.
2. Shopping App
- Created screens using Flutter and Firebase.
""", [
            {
                "title": "Resume Analyzer",
                "technologies": ["Python", "Flask"],
                "description": "Built an API using Python and Flask.",
            },
            {
                "title": "Shopping App",
                "technologies": ["Flutter", "Firebase"],
                "description": "Created screens using Flutter and Firebase.",
            },
        ]),
        ("aliases and duplicates", """Project: Web App
Tech Stack: python, JS, NodeJS, PYTHON
""", [{
            "title": "Web App",
            "technologies": ["Python", "JavaScript", "Node.js"],
            "description": None,
        }]),
        ("unknown explicitly listed technology", """Title: Research Tool
Tools: Python, QuantumKit
""", [{
            "title": "Research Tool",
            "technologies": ["Python", "QuantumKit"],
            "description": None,
        }]),
        ("wrapped title", """PROJECTS
AI-Based Deepfake Image Detection
Using Transfer Learning
- Tested image predictions.
""", [{
            "title": "AI-Based Deepfake Image Detection Using Transfer Learning",
            "technologies": [],
            "description": "Tested image predictions.",
        }]),
        (
            "inline technology list",
            "Resume Analyzer | Python, Flask",
            [{
                "title": "Resume Analyzer",
                "technologies": ["Python", "Flask"],
                "description": None,
            }],
        ),
        (
            "missing title",
            "Description: Built an API using Python.",
            [{
                "title": None,
                "technologies": ["Python"],
                "description": "Built an API using Python.",
            }],
        ),
        (
            "missing details",
            "Title: Library Management System",
            [{
                "title": "Library Management System",
                "technologies": [],
                "description": None,
            }],
        ),
        ("overlapping technology names", """Title: Mobile App
Technologies: React Native, ASP.NET Core
""", [{
            "title": "Mobile App",
            "technologies": ["React Native", "ASP.NET Core"],
            "description": None,
        }]),
        (
            "leading dot and explicit short language names",
            "Title: API Service\nTechnologies: .NET, C, R, Go",
            [{
                "title": "API Service",
                "technologies": [".NET", "C", "R", "Go"],
                "description": None,
            }],
        ),
        ("plain titles separated by descriptions", """PROJECTS
Weather Dashboard
Developed screens using React.

Sales Dashboard
Created reports using Python.
""", [
            {
                "title": "Weather Dashboard",
                "technologies": ["React"],
                "description": "Developed screens using React.",
            },
            {
                "title": "Sales Dashboard",
                "technologies": ["Python"],
                "description": "Created reports using Python.",
            },
        ]),
        (
            "decimal in description",
            "Title: Image Classifier\n- Achieved 95.5% accuracy.",
            [{
                "title": "Image Classifier",
                "technologies": [],
                "description": "Achieved 95.5% accuracy.",
            }],
        ),
        (
            "no project section",
            "SKILLS\nPython, Flask\nEDUCATION\nB.Tech 2027",
            [],
        ),
        ("no projects yet", "PROJECTS\nNo projects yet.", []),
        ("empty input", "", []),
    ]

    from schemas.resume_schema import Project

    for name, source, expected in cases:
        actual = extract_projects(source)
        assert actual == expected, (name, actual, expected)

        for entry in actual:
            assert Project(**entry).model_dump() == entry, entry

        print("PASS:", name)

    print("PASS: Project schema compatibility")
    print("\nEXPECTED JSON FOR THE REPORTED RESUME TEXT")
    print(json.dumps(
        extract_projects(REPORTED_TEXT),
        indent=4,
        ensure_ascii=False,
    ))


def check_file(path):
    if not path.is_file():
        raise FileNotFoundError(f"Resume not found: {path}")

    if path.suffix.lower() == ".pdf":
        from parser.pdf_parser import extract_text_from_pdf
        text = extract_text_from_pdf(str(path))

    elif path.suffix.lower() == ".docx":
        from docx import Document

        doc = Document(str(path))
        chunks = [paragraph.text for paragraph in doc.paragraphs]
        chunks.extend(
            " | ".join(cell.text for cell in row.cells)
            for table in doc.tables
            for row in table.rows
        )
        text = "\n".join(chunks)

    else:
        raise ValueError("Use a PDF or DOCX file.")

    if not text or not text.strip():
        raise ValueError(
            "No readable text. A scanned resume requires OCR."
        )

    print("\nRESUME PROJECTS JSON")
    print(json.dumps(
        extract_projects(text),
        indent=4,
        ensure_ascii=False,
    ))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("resume", nargs="?", help="Optional PDF/DOCX path")
    args = parser.parse_args()

    check_examples()

    if args.resume:
        check_file(Path(args.resume).expanduser().resolve())
    else:
        default = (
            Path(__file__).resolve().parent.parent
            / "sample_resume"
            / "Resume.pdf"
        )

        if default.is_file():
            check_file(default)
        else:
            print(
                '\nTo test your PDF: python test_projects.py '
                '"..\\sample_resume\\Resume.pdf"'
            )
