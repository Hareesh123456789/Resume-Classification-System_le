"""Test the experience extractor and optionally a PDF/DOCX resume."""

import argparse
import json
from pathlib import Path

from extractor.experience_extractor import extract_experience


def check_examples():
    cases = [
        (
            "internship",
            """INTERNSHIP EXPERIENCE
Python Full Stack Developer Intern - Example Technologies, Visakhapatnam
Nov 2023 - May 2024
- Developed web applications using Django and Flask.
- Integrated REST APIs.
PROJECT EXPERIENCE
AI-Based Image Detection
""",
            [{
                "company": "Example Technologies, Visakhapatnam",
                "role": "Python Full Stack Developer Intern",
                "duration": "Nov 2023 - May 2024",
                "description": (
                    "Developed web applications using Django and Flask.\n"
                    "Integrated REST APIs."
                ),
            }],
        ),
        (
            "multiple jobs",
            """WORK EXPERIENCE
Example Labs | Software Engineer | Jan 2022 - Present
Built payment APIs.
Data Analyst at Example Solutions
2020 - 2021
Created reporting dashboards.
EDUCATION
B.Tech 2020
""",
            [
                {
                    "company": "Example Labs",
                    "role": "Software Engineer",
                    "duration": "Jan 2022 - Present",
                    "description": "Built payment APIs.",
                },
                {
                    "company": "Example Solutions",
                    "role": "Data Analyst",
                    "duration": "2020 - 2021",
                    "description": "Created reporting dashboards.",
                },
            ],
        ),
        (
            "labeled entry",
            """Company: Example Company
Role: QA Specialist
Duration: 6 months
Description: Tested web and mobile applications.
""",
            [{
                "company": "Example Company",
                "role": "QA Specialist",
                "duration": "6 months",
                "description": "Tested web and mobile applications.",
            }],
        ),
        (
            "separate lines",
            """Software Engineer
Example Company Ltd
June 2022 - August 2023
Maintained APIs.
""",
            [{
                "company": "Example Company Ltd",
                "role": "Software Engineer",
                "duration": "June 2022 - August 2023",
                "description": "Maintained APIs.",
            }],
        ),
        (
            "missing fields",
            "Role: Research Fellow",
            [{
                "company": None,
                "role": "Research Fellow",
                "duration": None,
                "description": None,
            }],
        ),
        (
            "project-only resume",
            """CAREER OBJECTIVE
Seeking an internship opportunity in AI and ML.
SKILLS
Python, HTML, Oracle
EDUCATION
B.Tech. (AI&ML)
2027
PROJECT EXPERIENCE
AI-Based Deepfake Image Detection Using Transfer Learning
CERTIFICATIONS
Introduction to Machine Learning
DECLARATION
The above information is correct.
""",
            [],
        ),
        (
            "no experience",
            "EXPERIENCE\nNo professional experience yet.",
            [],
        ),
        ("empty input", "", []),
        (
            "bulleted description",
            """WORK EXPERIENCE
Developer - Example Company
2022 - 2024
- Software Engineer mentoring and training.
""",
            [{
                "company": "Example Company",
                "role": "Developer",
                "duration": "2022 - 2024",
                "description": "Software Engineer mentoring and training.",
            }],
        ),
        (
            "ISO dates",
            "Intern - Example Company\n2023-06 - 2023-08",
            [{
                "company": "Example Company",
                "role": "Intern",
                "duration": "2023-06 - 2023-08",
                "description": None,
            }],
        ),
        (
            "bulleted job header",
            """WORK EXPERIENCE
- Software Engineer - Example Company
2022 - 2024
- Debugged APIs.
""",
            [{
                "company": "Example Company",
                "role": "Software Engineer",
                "duration": "2022 - 2024",
                "description": "Debugged APIs.",
            }],
        ),
    ]

    from schemas.resume_schema import Experience

    for name, source, expected in cases:
        actual = extract_experience(source)
        assert actual == expected, (name, actual, expected)

        for entry in actual:
            assert Experience(**entry).model_dump() == entry, entry

        print("PASS:", name)

    print("PASS: Experience schema compatibility")
    print("\nSAMPLE EXPERIENCE JSON")
    print(json.dumps(
        extract_experience(cases[0][1]),
        indent=4,
        ensure_ascii=False,
    ))


def check_file(path):
    if not path.is_file():
        raise FileNotFoundError(f"Resume not found: {path}")

    if path.suffix.lower() == ".pdf":
        from parser.pdf_parser import extract_text_from_pdf
        text = extract_text_from_pdf(str(path))

    elif path.suffix.lower() == ".docx":
        from docx import Document

        doc = Document(str(path))
        chunks = [paragraph.text for paragraph in doc.paragraphs]
        chunks.extend(
            " | ".join(cell.text for cell in row.cells)
            for table in doc.tables
            for row in table.rows
        )
        text = "\n".join(chunks)

    else:
        raise ValueError("Use a PDF or DOCX file.")

    if not text or not text.strip():
        raise ValueError("No readable text. A scanned resume requires OCR.")

    print("\nRESUME EXPERIENCE JSON")
    print(json.dumps(
        extract_experience(text),
        indent=4,
        ensure_ascii=False,
    ))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("resume", nargs="?", help="Optional PDF/DOCX path")
    args = parser.parse_args()

    check_examples()

    if args.resume:
        check_file(Path(args.resume).expanduser().resolve())
    else:
        default = (
            Path(__file__).resolve().parent.parent
            / "sample_resume"
            / "Resume.pdf"
        )

        if default.is_file():
            check_file(default)
        else:
            print(
                '\nTo test your PDF: python test_experience.py '
                '"..\\sample_resume\\Resume.pdf"'
            )
