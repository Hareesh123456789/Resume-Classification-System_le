"""Check certification extraction, then optionally analyze a PDF/DOCX."""

import argparse
import json
from pathlib import Path

from extractor.certification_extractor import extract_certifications


REPORTED_TEXT = """EDUCATION
B.Tech. (AI&ML)
2027
PROJECT EXPERIENCE
AI-Based Deepfake Image Detection Using Transfer Learning
1. Deepfake Image Detection.
2.High Detection Accuracy.
CERTIFICATIONS
1) Introduction to Machine learning (12 weeks)
2) Introduction to Artificial intelligence  (12 weeks)
Aditya institute of technology and
management, Tekkali
3.Fast and Efficient Prediction.
DECLARATION
The above information is correct.
"""

REPORTED_EXPECTED = [
    {
        "name": "Introduction to Machine learning",
        "organization": None,
        "year": None,
    },
    {
        "name": "Introduction to Artificial intelligence",
        "organization": None,
        "year": None,
    },
]


def check_examples():
    cases = [
        (
            "reported PDF text, durations, displaced text",
            REPORTED_TEXT,
            REPORTED_EXPECTED,
        ),
        ("multiple certificates with separators", """CERTIFICATIONS
AWS Certified Cloud Practitioner | Amazon Web Services | 2024
Python for Everybody - Coursera - 2023
PROJECTS
Python App
""", [
            {"name": "AWS Certified Cloud Practitioner",
             "organization": "Amazon Web Services", "year": "2024"},
            {"name": "Python for Everybody",
             "organization": "Coursera", "year": "2023"},
        ]),
        ("explicit labels", """Name: Communication Skills
Organization: Example Learning Academy
Year: 2022
Certificate: Research Methods
Issuer: Example Council
Issue Date: June 2024
""", [
            {"name": "Communication Skills",
             "organization": "Example Learning Academy", "year": "2022"},
            {"name": "Research Methods",
             "organization": "Example Council", "year": "2024"},
        ]),
        (
            "inline provider",
            "1. Python for Everybody from Coursera (2023)",
            [{"name": "Python for Everybody",
              "organization": "Coursera", "year": "2023"}],
        ),
        (
            "generic explicit issuer",
            "Certificate: Safety Awareness issued by Example Board (2024)",
            [{"name": "Safety Awareness",
              "organization": "Example Board", "year": "2024"}],
        ),
        (
            "provider on separate line",
            "CERTIFICATIONS\nPython Basics\nCoursera\n2024",
            [{"name": "Python Basics",
              "organization": "Coursera", "year": "2024"}],
        ),
        ("wrapped certificate and issuer", """Certificate: Introduction to Artificial
Intelligence (12 weeks)
Provider: Example Institute of
Technology
Completed: July 2024
""", [
            {"name": "Introduction to Artificial Intelligence",
             "organization": "Example Institute of Technology", "year": "2024"},
        ]),
        (
            "missing issuer and year",
            "CERTIFICATIONS\n1) Communication Skills",
            [{"name": "Communication Skills",
              "organization": None, "year": None}],
        ),
        (
            "missing certificate name",
            "Name:\nProvider: IBM\nYear: 2024",
            [{"name": None, "organization": "IBM", "year": "2024"}],
        ),
        (
            "duration is not a year",
            "CERTIFICATIONS\nIntroduction to Python (12 weeks)",
            [{"name": "Introduction to Python",
              "organization": None, "year": None}],
        ),
        (
            "issue year and expiry year",
            "Name: Cloud Fundamentals\nIssued: May 2023\nValid until: May 2026",
            [{"name": "Cloud Fundamentals",
              "organization": None, "year": "2023"}],
        ),
        (
            "expiry alone is not completion year",
            "Name: Cloud Fundamentals\nExpires: 2026",
            [{"name": "Cloud Fundamentals",
              "organization": None, "year": None}],
        ),
        (
            "ambiguous years",
            "Name: Cloud Fundamentals\n2022\n2023",
            [{"name": "Cloud Fundamentals",
              "organization": None, "year": None}],
        ),
        (
            "numeric issue date and expiry",
            "Name: Cloud Fundamentals\nIssued: 26/05/2024; Expires: 26/05/2026",
            [{"name": "Cloud Fundamentals",
              "organization": None, "year": "2024"}],
        ),
        (
            "wrapped labeled institution",
            "Name: Research Methods\nProvider: Example\nUniversity",
            [{"name": "Research Methods",
              "organization": "Example University", "year": None}],
        ),
        (
            "certificate name beginning with a dot",
            "Certificate: .NET Fundamentals",
            [{"name": ".NET Fundamentals",
              "organization": None, "year": None}],
        ),
        (
            "table header",
            "CERTIFICATIONS\nName | Issuer | Year\nPython Basics | Coursera | 2024",
            [{"name": "Python Basics",
              "organization": "Coursera", "year": "2024"}],
        ),
        (
            "duplicate certificates",
            "CERTIFICATIONS\n1) Python Basics\n2) Python Basics",
            [{"name": "Python Basics",
              "organization": None, "year": None}],
        ),
        (
            "no certification section",
            "SKILLS\nPython\nEDUCATION\nB.Tech 2027",
            [],
        ),
        (
            "no certificates yet",
            "CERTIFICATIONS\nNo certifications yet.",
            [],
        ),
        ("empty input", "", []),
    ]

    from schemas.resume_schema import Certification

    for name, source, expected in cases:
        actual = extract_certifications(source)
        assert actual == expected, (name, actual, expected)

        for entry in actual:
            assert Certification(**entry).model_dump() == entry, entry

        print("PASS:", name)

    print("PASS: Certification schema compatibility")
    print("\nEXPECTED JSON FOR THE REPORTED RESUME TEXT")
    print(json.dumps(
        extract_certifications(REPORTED_TEXT),
        indent=4,
        ensure_ascii=False,
    ))


def check_file(path):
    if not path.is_file():
        raise FileNotFoundError(f"Resume not found: {path}")

    if path.suffix.lower() == ".pdf":
        from parser.pdf_parser import extract_text_from_pdf
        text = extract_text_from_pdf(str(path))

    elif path.suffix.lower() == ".docx":
        from docx import Document

        doc = Document(str(path))
        chunks = [paragraph.text for paragraph in doc.paragraphs]
        chunks.extend(
            " | ".join(cell.text for cell in row.cells)
            for table in doc.tables
            for row in table.rows
        )
        text = "\n".join(chunks)

    else:
        raise ValueError("Use a PDF or DOCX file.")

    if not text or not text.strip():
        raise ValueError(
            "No readable text. A scanned resume requires OCR."
        )

    print("\nRESUME CERTIFICATIONS JSON")
    print(json.dumps(
        extract_certifications(text),
        indent=4,
        ensure_ascii=False,
    ))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("resume", nargs="?", help="Optional PDF/DOCX path")
    args = parser.parse_args()

    check_examples()

    if args.resume:
        check_file(Path(args.resume).expanduser().resolve())
    else:
        default = (
            Path(__file__).resolve().parent.parent
            / "sample_resume"
            / "Resume.pdf"
        )

        if default.is_file():
            check_file(default)
        else:
            print(
                '\nTo test your PDF: python test_certifications.py '
                '"..\\sample_resume\\Resume.pdf"'
            )
