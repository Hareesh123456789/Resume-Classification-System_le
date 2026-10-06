r"""Run education checks, then optionally analyze one PDF/DOCX.

python test_education.py
python test_education.py "..\sample_resume\Resume.pdf"
"""

import argparse
import json
from pathlib import Path

from extractor.education_extractor import extract_education


def check_reported_pdf_text():
    # Regression: preserve the exact line breaks and fragmented tokens reported
    # by the PDF parser. These are test inputs, not extraction rules.
    text = """SKILLS
Python, Oracle
EDUCATION
Academic
Qualification
University/Board
Year
Percentage
/CGPA
B.Tech. (AI&ML)
Aditya institute of technology and
management, Tekkali
2027
7.70 CGPA
D iploma in E CE
2024
7.7
SSC
Sri Gnana J yothi school, N.peta
2021
9.7
PROJECT EXPERIENCE
AI-Based Deepfake Image Detection Using Transfer Learning
CERTIFICATIONS
Aditya institute of technology and
management, Tekkali
"""
    actual = extract_education(text)
    expected = [
        {"degree": "B.Tech", "specialization": "AI&ML",
         "institution": "Aditya institute of technology and management, Tekkali",
         "university": None, "year": "2027", "score": "7.70 (CGPA)"},
        {"degree": "Diploma", "specialization": "ECE", "institution": None,
         "university": None, "year": "2024", "score": "7.7"},
        {"degree": "Class 10", "specialization": None,
         "institution": "Sri Gnana J yothi school, N.peta",
         "university": None, "year": "2021", "score": "9.7"},
    ]
    assert actual == expected, actual
    wrapped = extract_education("Diploma in E CE\nGovernment\nPolytechnic College,\nSample City\nSBTET-AP\n2024\n66%")
    assert wrapped[0]["institution"] == "Government Polytechnic College, Sample City", wrapped
    assert wrapped[0]["university"] == "SBTET-AP", wrapped
    assert extract_education("B.Tech\n2020\n2024")[0]["year"] is None
    assert extract_education("SSC\n2021\n9.7\n8.8")[0]["score"] is None
    print("PASS: reported PDF text, fragmented Diploma/ECE, wrapped names, record boundaries, raw scores")


def check_examples():
    text = """EDUCATION
B.Tech in Computer Science and Engineering
Example Institute of Technology
Example Technical University
2021 - 2025 | CGPA: 8.2/10
Diploma in Computer Engineering
Example Polytechnic College
State Board of Technical Education
2018 - 2021 | 78%
SSC | Example High School | CBSE | 2018 | 91%
PROJECTS
Graduation prediction project: Python, 2026
"""
    records = extract_education(text)
    assert len(records) == 3, records
    assert records[0]["degree"] == "B.Tech", records
    assert records[0]["specialization"] == "Computer Science and Engineering", records
    assert records[0]["institution"] == "Example Institute of Technology", records
    assert records[0]["university"] == "Example Technical University", records
    assert records[0]["year"] == "2025", records
    assert records[0]["score"] == "8.2/10 (CGPA)", records
    assert records[1]["year"] == "2021", records
    assert records[1]["score"] == "78%", records
    assert records[2]["score"] == "91%", records
    assert records[2]["university"] == "CBSE", records
    assert extract_education("Python, SQL, Docker") == []
    assert extract_education("") == []
    ongoing = extract_education("M.Sc in Physics\n2024 - Present")
    assert ongoing[0]["year"] is None, ongoing
    missing = extract_education("BCA")
    assert all(value is None for key, value in missing[0].items() if key != "degree"), missing
    assert extract_education("B.Sc | 2022 | 8.6(CGPA)")[0]["score"] == "8.6 (CGPA)"
    assert extract_education("B.Tech | CGPA: 80")[0]["score"] is None
    assert extract_education("B.Tech | GPA: 8/4")[0]["score"] is None
    assert extract_education("To be a skilled engineer") == []
    assert extract_education("BE in Electronics\n2020\nGPA: 3.8/4")[0]["score"] == "3.8/4 (GPA)"
    assert extract_education("B.Tech\nExpected: 2027\n2023 - Present")[0]["year"] == "2027"
    rows = extract_education("B.Tech | CSE | Test College | 2024 | 86%\nMBA | Test University | 2026 | GPA: 3.4/4")
    assert len(rows) == 2 and rows[0]["year"] == "2024" and rows[1]["year"] == "2026", rows
    assert rows[0]["specialization"] == "CSE", rows
    # Independent calls must not retain records from the previous resume.
    assert extract_education("Python and SQL") == []
    print("PASS: degree, specialization, institutions, years, scores, missing fields, section boundary")
    # Validate against your existing schema when installed in the backend.
    from schemas.resume_schema import Education
    for record in records:
        validated = Education(**record)
        assert validated.model_dump() == record, validated
    print("PASS: compatibility with the Education JSON schema")
    print(json.dumps(records, indent=4, ensure_ascii=False))


def check_file(path):
    if not path.is_file():
        raise FileNotFoundError(f"Resume not found: {path}")
    if path.suffix.lower() == ".pdf":
        from parser.pdf_parser import extract_text_from_pdf
        text = extract_text_from_pdf(str(path))
    elif path.suffix.lower() == ".docx":
        from docx import Document
        doc = Document(str(path))
        chunks = [p.text for p in doc.paragraphs]
        chunks.extend(" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
        text = "\n".join(chunks)
    else:
        raise ValueError("Use a PDF or DOCX file.")
    if not text or not text.strip():
        raise ValueError("No readable text. A scanned resume requires OCR.")
    print("\nRESUME EDUCATION JSON")
    print(json.dumps(extract_education(text), indent=4, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("resume", nargs="?", help="Optional PDF/DOCX path")
    args = parser.parse_args()
    check_examples()
    check_reported_pdf_text()
    if args.resume:
        check_file(Path(args.resume).expanduser().resolve())
    else:
        default = Path(__file__).resolve().parent.parent / "sample_resume" / "Resume.pdf"
        if default.is_file():
            check_file(default)
        else:
            print("\nTo test your PDF: python test_education.py \"..\\sample_resume\\Resume.pdf\"")
